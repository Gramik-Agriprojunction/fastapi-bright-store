"""Equivalent questions, filters, duplicates, RTO receipt status, and Excel."""

from datetime import date, datetime

from openpyxl import load_workbook

from app.modules.webtechAIAgent.reconciliation import excel_mail, report_generation
from app.modules.webtechAIAgent.reconciliation.entity_normalization import search_terms
from app.modules.webtechAIAgent.reconciliation.query_execution import execute_report
from app.modules.webtechAIAgent.reconciliation.question_understanding import understand
from app.modules.webtechAIAgent.reconciliation.service import _email_report, ask_reconciliation
from app.modules.webtechAIAgent.reconciliation.workbook import Book, Move
from tests.test_reconciliation_qa import TODAY, _patch

DAY = date(2026, 10, 15)


def _move(
    day: str,
    bucket: str,
    *,
    doc: str,
    qty_in: float | None = None,
    qty_out: float | None = None,
    sale: str = "",
    state: str = "done",
    demand: float | None = None,
    dest: str = "",
) -> Move:
    return Move(
        when=datetime.fromisoformat(day),
        doc=doc,
        bucket=bucket,
        state=state,
        sale=sale,
        partner="Ram",
        source="Hub",
        dest=dest,
        qty_in=qty_in,
        qty_out=qty_out,
        running=None,
        return_of=sale,
        lot="",
        note="",
        demand=demand,
    )


def _book(moves: list[Move], *, on_hand: float = 2, name: str = "Musturd 1666", sku: str = "1445") -> Book:
    return Book(
        path="odoo",
        product_name=name,
        sku=sku,
        lot="",
        store_name="Lavkush Dark Store",
        warehouse="Lavkush Dark Store",
        location="BL-D/Stock",
        snapshot=DAY,
        company="",
        quant_id="",
        live_on_hand=on_hand,
        live_reserved=1,
        live_free=on_hand - 1,
        transit_on_hand=0,
        moves=moves,
    )


def test_equivalent_questions_share_an_intent() -> None:
    groups = {
        "SALES_LEDGER": (
            "Pure sale ka ledger chahiye.",
            "mujhe sales ledger do",
            "sale ledgar chahiye",
        ),
        "SALES_OUT_REPORT": (
            "Mujhe product-wise sales details do.",
            "product wise sale nikalo",
        ),
        "DELIVERED_ORDERS": (
            "Delivered orders ka order date aur quantity batao.",
            "delivered order list do",
        ),
        "RTO_RECEIVED_REPORT": (
            "Jo RTO dark store ko receive ho chuke hain unki details do.",
            "RTO received report nikalo",
        ),
        "OUTGOING_STOCK_REPORT": (
            "Kaun se orders abhi outgoing hain?",
            "Jo stock store se jane wala hai uski list do.",
        ),
    }
    for intent, questions in groups.items():
        found = {understand(question, DAY)["intent"] for question in questions}
        assert found == {intent}

    excel = understand("Excel bana do, same format mein jaise previous report bani thi.", DAY)
    mail = understand("Is report ko is email par bhej do: user@example.com.", DAY)
    assert excel["delivery"] == "GENERATE_EXCEL_REPORT"
    assert excel["intent"] is None
    assert mail["delivery"] == "EMAIL_REPORT"
    assert "user" not in mail["terms"]
    assert "example" not in mail["terms"]


def test_clear_question_does_not_offer_suggested_questions(monkeypatch) -> None:
    _patch(monkeypatch)
    answer = ask_reconciliation("lavkush mustard kitna hai", today=DAY)
    assert answer["suggestions"] == []
    assert "Kya chahiye" not in answer["reply"]

    ledger = ask_reconciliation("Pure sale ka ledger chahiye.", today=DAY)
    assert "Kaun sa warehouse? Naam Select Karo." in ledger["reply"]
    assert "Lavkush Dark Store Barabanki (LAVK)" in ledger["suggestions"]
    assert "on hand, sale out" not in ledger["reply"]


def test_named_columns_and_duplicate_orders_are_counted_once() -> None:
    parsed = understand("Order number, sale out aur RTO IN ki list nikalo.", DAY)
    assert parsed["columns"] == ["order_number", "sale_out", "rto_in"]
    book = _book(
        [
            _move("2026-09-23", "Delivered", doc="OUT/1", qty_out=10, sale="AGRI1"),
            _move("2026-09-24", "Delivered", doc="OUT/2", qty_out=5, sale="AGRI1"),
            _move("2026-09-23", "Delivered", doc="OUT/1", qty_out=10, sale="AGRI1"),
            _move("2026-10-08", "RTO", doc="RTO/1", qty_in=4, sale="AGRI1"),
            _move("2026-10-09", "RTO in transit", doc="RTO/2", qty_in=9, sale="AGRI1", state="assigned"),
        ]
    )
    result = execute_report(
        book,
        {"intent": "SALES_LEDGER", "message": "Order number, sale out aur RTO IN ki list nikalo.", "columns": parsed["columns"], "date_filter": {"type": "all"}},
        DAY,
    )
    assert result["record_count"] == 1
    assert "Sale Out total 15" in result["reply"]
    assert "RTO IN total 4" in result["reply"]
    assert result["table"]["record_count"] == 1
    assert result["table"]["rows"][0]["Sale Out"] == "15"
    assert result["table"]["rows"][0]["RTO IN"] == "4"
    assert "RTO/2" not in str(result["table"])
    assert "Orders 1" in result["reply"]


def test_rto_received_excludes_transit_and_open_returns() -> None:
    book = _book(
        [
            _move("2026-10-08", "RTO", doc="RTO/1", qty_in=4, sale="AGRI1"),
            _move("2026-10-09", "RTO", doc="RTO/OPEN", qty_in=6, sale="AGRI2", state="assigned"),
            _move("2026-10-10", "RTO in transit", doc="RTO/T", qty_in=8, sale="AGRI3", state="assigned"),
        ]
    )
    result = execute_report(
        book,
        {"intent": "RTO_RECEIVED_REPORT", "message": "RTO receive ho chuke", "date_filter": {"type": "all"}},
        DAY,
    )
    assert result["table"]["title"] == "RTO"
    assert result["table"]["columns"] == ["Received Qty"]
    assert result["table"]["rows"][0]["Received Qty"] == "4"
    assert result["table"]["record_count"] == 1
    assert "Received Qty total 4" in result["reply"]
    assert "RTO/OPEN" not in str(result["table"])
    assert "RTO/T" not in str(result["table"])
    assert "Order Date is not stored" in result["reply"]


def test_outgoing_statuses_stay_separate() -> None:
    book = _book(
        [
            _move("2026-10-11", "Outgoing", doc="OUT/A", qty_out=3, sale="AGRI1", state="assigned", dest="Customer A"),
            _move("2026-10-11", "Outgoing", doc="OUT/P", qty_out=2, sale="AGRI2", state="confirmed", dest="Customer B"),
            _move("2026-10-12", "Cancelled", doc="OUT/C", sale="AGRI3", state="cancel", demand=8),
            _move("2026-10-01", "Delivered", doc="OUT/D", qty_out=5, sale="AGRI4"),
        ]
    )
    result = execute_report(
        book,
        {
            "intent": "OUTGOING_STOCK_REPORT",
            "message": "Kaun se orders abhi outgoing hain?",
            "date_filter": {"type": "all"},
        },
        DAY,
    )
    assert "assigned: 1 orders · 3 qty" in result["reply"]
    assert "pending: 1 orders · 2 qty" in result["reply"]
    assert "OUT/C" not in result["reply"]
    assert "OUT/D" not in result["reply"]
    assert "8" not in result["reply"]


def test_delivered_date_filter_and_empty_result() -> None:
    book = _book(
        [
            _move("2026-09-23", "Delivered", doc="OUT/1", qty_out=21, sale="AGRI1"),
            _move("2026-10-01", "Delivered", doc="OUT/2", qty_out=17, sale="AGRI2"),
        ]
    )
    ranged = execute_report(
        book,
        {
            "intent": "DELIVERED_ORDERS",
            "message": "delivered orders",
            "date_filter": {"type": "range", "start": "2026-10-01", "end": "2026-10-15"},
        },
        DAY,
    )
    assert ranged["table"]["rows"][0]["Order Number"] == "AGRI2"
    assert ranged["table"]["record_count"] == 1
    assert "Delivered Qty total 17" in ranged["reply"]

    empty = execute_report(
        _book([_move("2026-09-19", "IN", doc="IN/1", qty_in=36)]),
        {"intent": "DELIVERED_ORDERS", "message": "delivered orders", "date_filter": {"type": "all"}},
        DAY,
    )
    assert empty["rows"] == []
    assert "No matching records found" in empty["reply"]


def test_two_products_keep_their_own_totals() -> None:
    mustard = execute_report(
        _book([_move("2026-10-01", "Delivered", doc="OUT/M", qty_out=10, sale="AGRI1")]),
        {"intent": "SALES_OUT_REPORT", "message": "product wise sales", "date_filter": {"type": "all"}},
        DAY,
    )
    gramino = execute_report(
        _book(
            [_move("2026-10-02", "Delivered", doc="OUT/G", qty_out=4, sale="AGRI9")],
            name="Gramino",
            sku="200",
        ),
        {"intent": "SALES_OUT_REPORT", "message": "product wise sales", "date_filter": {"type": "all"}},
        DAY,
    )
    assert "Musturd 1666" in mustard["reply"]
    assert "Sale Out total 10" in mustard["reply"]
    assert "Gramino" in gramino["reply"]
    assert "Sale Out total 4" in gramino["reply"]
    assert "Sale Out total 14" not in mustard["reply"]
    assert "Sale Out total 14" not in gramino["reply"]


def test_on_hand_is_the_warehouse_location_quantity() -> None:
    book = _book(
        [
            _move("2026-09-19", "IN", doc="IN/1", qty_in=36),
            _move("2026-09-23", "Delivered", doc="OUT/1", qty_out=21, sale="AGRI1"),
            _move("2026-10-01", "Delivered", doc="OUT/2", qty_out=17, sale="AGRI2"),
            _move("2026-10-08", "RTO", doc="RTO/1", qty_in=4, sale="AGRI1"),
        ],
        on_hand=7,
    )
    result = execute_report(
        book,
        {"intent": "PRODUCT_STOCK_SUMMARY", "message": "product stock summary", "date_filter": {"type": "current"}},
        DAY,
    )
    assert "On Hand Qty: 7" in result["reply"]
    assert "Expected Qty: 2" in result["reply"]
    assert "Variance: 5" in result["reply"]
    assert "physical stock at BL-D/Stock" in result["reply"]


def test_excel_uses_filtered_rows_and_does_not_invent_orders(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(report_generation, "OUT_DIR", tmp_path)
    book = _book(
        [
            _move("2026-09-23", "Delivered", doc="OUT/1", qty_out=10, sale="AGRI1"),
            _move("2026-09-24", "Delivered", doc="OUT/2", qty_out=5, sale="AGRI1"),
            _move("2026-09-23", "Delivered", doc="OUT/1", qty_out=10, sale="AGRI1"),
            _move("2026-10-08", "RTO", doc="RTO/1", qty_in=4, sale="AGRI1"),
        ]
    )
    saved = report_generation.save_report(
        [book],
        {
            "intent": "SALES_LEDGER",
            "message": "pure sale ka ledger",
            "date_filter": {"type": "all"},
        },
        DAY,
    )
    assert not isinstance(saved, str)
    workbook = load_workbook(saved["path"])
    assert workbook.sheetnames == ["Summary", "Sales Ledger"]
    assert workbook["Summary"]["B3"].value == 1
    assert workbook["Summary"]["A4"].value == "Sale Out"
    assert workbook["Summary"]["B4"].value == 15
    assert workbook["Summary"]["B5"].value == 4
    ledger = workbook["Sales Ledger"]
    assert ledger["A1"].value == "Order Number"
    assert ledger.max_row == 3
    assert ledger["A2"].value == "AGRI1"
    assert ledger.auto_filter.ref.startswith("A1:")
    workbook.close()


def test_two_warehouses_and_a_date_filter(monkeypatch) -> None:
    _patch(monkeypatch)
    both = ask_reconciliation("fatehpur aur lavkush mustard delivered orders", today=DAY)
    assert "Fatehpur Dark Store" in both["reply"]
    assert "Lavkush Dark Store" in both["reply"]
    assert both["suggestions"] == []

    ranged = ask_reconciliation(
        "lavkush mustard delivered orders 1 october 2026 se 15 october 2026",
        today=DAY,
    )
    orders = [row.get("Order Number") for table in ranged["tables"] for row in table["rows"]]
    assert "AGRI2" in orders
    assert "AGRI1" not in orders
    assert ranged["display_mode"] == "chat"
    assert ranged["files"] == []
    assert ranged["suggestions"] == []


def test_email_is_not_claimed_when_sending_fails(monkeypatch, tmp_path) -> None:
    _patch(monkeypatch)
    monkeypatch.setattr(report_generation, "OUT_DIR", tmp_path)
    monkeypatch.setattr(excel_mail, "OUT_DIR", tmp_path)

    class Settings:
        ses_smtp_host = ""
        ses_smtp_port = 587
        ses_smtp_user = ""
        ses_smtp_pass = ""
        ses_from_email = "noreply@gramik.in"
        ses_from_name = "WebTechAgent AI Assistant"

    monkeypatch.setattr(excel_mail, "get_settings", lambda: Settings())
    answer = ask_reconciliation(
        "lavkush mustard delivered orders excel banao aur user@example.com par bhej do",
        today=DAY,
    )
    assert "bhej diya" not in answer["reply"]
    assert "Email nahi gaya" in answer["reply"] or "configured nahi" in answer["reply"]
    assert answer["files"]
    assert "Download:" in answer["reply"]


def test_empty_report_is_not_emailed(monkeypatch) -> None:
    called = {"send": False}

    def fail_send(recipients: list[str], paths: list[str]) -> str:
        called["send"] = True
        return "Excel bhej diya: user@example.com."

    monkeypatch.setattr(
        "app.modules.webtechAIAgent.reconciliation.service.send_excel_files",
        fail_send,
    )
    reply = _email_report("user@example.com par bhej do", ["/tmp/missing.xlsx"], 0)
    assert called["send"] is False
    assert "bhej diya" not in reply
    assert "records nahi" in reply


def test_ramnagar_stays_a_warehouse_term() -> None:
    assert "ramnagar" in search_terms("Ram Nagar mustard kitna hai")
    assert "musturd" in search_terms("Ram Nagar mustard kitna hai")


def test_display_mode_is_chat_unless_excel_is_explicit() -> None:
    chat_questions = (
        "Fatehpur store ka Sale Out batao",
        "Delivered orders ki list dikhao",
        "RTO IN ka total kitna hai?",
        "Pure sale ka ledger dikhao",
    )
    excel_questions = (
        "Pure sale ka ledger Excel mein bana do",
        "Is report ka Excel download karna hai",
        "Delivered orders ki details Excel mein do",
        "report bana kar download kar do",
        "export to xlsx",
    )
    for question in chat_questions:
        parsed = understand(question, DAY)
        assert parsed["display_mode"] == "chat", question
    for question in excel_questions:
        parsed = understand(question, DAY)
        assert parsed["display_mode"] == "excel", question
    assert understand("RTO IN ka total kitna hai?", DAY)["intent"] == "RTO_RECEIVED_REPORT"
    assert understand("Fatehpur store ka Sale Out batao", DAY)["intent"] == "SALES_OUT_REPORT"


def test_chat_preview_does_not_create_excel_and_excel_uses_the_same_rows(monkeypatch, tmp_path) -> None:
    _patch(monkeypatch)
    monkeypatch.setattr(report_generation, "OUT_DIR", tmp_path)
    chat = ask_reconciliation("lavkush mustard delivered orders ki list dikhao", today=DAY)
    assert chat["display_mode"] == "chat"
    assert chat["files"] == []
    assert "Excel:" not in chat["reply"]
    assert chat["table"]["record_count"] == 2
    assert "Showing 2 of 2 records." in chat["reply"]

    exported = ask_reconciliation("lavkush mustard delivered orders ki details Excel mein do", today=DAY)
    assert exported["display_mode"] == "excel"
    assert exported["files"]
    assert exported["table"]["record_count"] == chat["table"]["record_count"]
    workbook = load_workbook(exported["files"][0]["path"])
    assert workbook["Delivered Orders"].max_row == 4
    workbook.close()


def test_large_result_shows_a_page_and_excel_keeps_every_row(tmp_path) -> None:
    moves = [
        _move("2026-10-01", "Delivered", doc=f"OUT/{index}", qty_out=1, sale=f"AGRI{index}")
        for index in range(1, 26)
    ]
    book = _book(moves)
    first = execute_report(
        book,
        {"intent": "DELIVERED_ORDERS", "message": "delivered orders", "date_filter": {"type": "all"}, "page": 1},
        DAY,
    )
    second = execute_report(
        book,
        {"intent": "DELIVERED_ORDERS", "message": "aur dikhao", "date_filter": {"type": "all"}, "page": 2},
        DAY,
    )
    assert first["table"]["record_count"] == 25
    assert len(first["table"]["rows"]) == 20
    assert first["table"]["rows"][0]["Order Number"] == "AGRI25"
    assert "Showing 1-20 of 25 records." in first["reply"]
    assert "Excel" not in first["reply"]
    assert len(second["table"]["rows"]) == 5
    assert len(second["rows"]) == 25
    saved = report_generation.save_report(
        [book],
        {"intent": "DELIVERED_ORDERS", "message": "delivered orders", "date_filter": {"type": "all"}, "page": 1},
        DAY,
        dest=tmp_path / "delivered.xlsx",
    )
    assert not isinstance(saved, str)
    workbook = load_workbook(saved["path"])
    assert workbook["Delivered Orders"].max_row == 27
    workbook.close()
