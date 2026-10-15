from datetime import date, datetime

from fastapi.testclient import TestClient

from app.main import app
from app.modules.webtechAIAgent.reconciliation import excel_export, odoo_ledger
from app.modules.webtechAIAgent.reconciliation.question_understanding import understand
from app.modules.webtechAIAgent.reconciliation.service import answer_question, ask_reconciliation
from app.modules.webtechAIAgent.reconciliation.workbook import Book, Move
from app.modules.webtechAIAgent.service import MoveRow, ProductRow, Warehouse

client = TestClient(app)
TODAY = date(2026, 10, 15)


def _move(day: str, bucket: str, qty_in: float | None = None, qty_out: float | None = None) -> Move:
    return Move(
        when=datetime.fromisoformat(day),
        doc=f"{bucket}-{day}",
        bucket=bucket,
        state="done",
        sale="",
        partner="",
        source="",
        dest="",
        qty_in=qty_in,
        qty_out=qty_out,
        running=None,
        return_of="",
        lot="",
        note="",
    )


def _book() -> Book:
    return Book(
        path="odoo",
        product_name="Musturd 1666",
        sku="1445",
        lot="",
        store_name="Lavkush Dark Store",
        warehouse="Lavkush Dark Store",
        location="BL-D/Stock",
        snapshot=TODAY,
        company="",
        quant_id="",
        live_on_hand=2,
        live_reserved=2,
        live_free=0,
        transit_on_hand=0,
        moves=[
            _move("2026-09-19", "IN", qty_in=36),
            _move("2026-09-23", "Delivered", qty_out=21),
            _move("2026-10-01", "Delivered", qty_out=17),
            _move("2026-10-08", "RTO", qty_in=4),
        ],
    )


def test_same_formula_as_the_reconciliation_sheet() -> None:
    book = _book()
    gap = answer_question(book, {"intent": "stock_variance", "date_filter": {"type": "current"}, "mode": "total"}, TODAY)
    sales = answer_question(
        book,
        {"intent": "date_range_sales", "date_filter": {"type": "all"}, "mode": "total"},
        TODAY,
    )
    assert "Hona chahiye 2" in gap["reply"]
    assert "Live on hand 2" in gap["reply"]
    assert "Gap 0" in gap["reply"]
    assert "Total 38 qty" in sales["reply"]


def test_sara_kuch_means_on_hand_sale_rto_in_dc_and_back() -> None:
    parsed = understand("sara kuch de do", TODAY)
    assert parsed["intent"] == "everything"
    assert parsed["handled_signal"] is True
    assert parsed["terms"] == []

    book = _book()
    book.moves.append(_move("2026-09-20", "IN", qty_in=10))
    book.moves[-1].source = "RC Lucknow (DC-01)/Stock"
    book.moves.append(_move("2026-09-21", "Other OUT", qty_out=3))
    reply = answer_question(
        book,
        {"intent": "everything", "date_filter": {"type": "all"}, "mode": "total"},
        TODAY,
    )["reply"]
    assert "On hand 2" in reply
    assert "Sale out 38" in reply
    assert "RTO 4" in reply
    assert "In 46" in reply
    assert "1. IN-2026-09-19 · 36" in reply
    assert "2. IN-2026-09-20 · 10" in reply
    assert "DC In 38" in reply
    assert "DC count 2" in reply
    assert "1. Delivered-2026-09-23 · 21" in reply
    assert "2. Delivered-2026-10-01 · 17" in reply
    assert "Back 4" in reply
    assert "1. RTO-2026-10-08 · 4" in reply

    assert understand("in", TODAY)["intent"] == "stock_inward"
    assert understand("dc in", TODAY)["intent"] == "delivered_dc"
    assert understand("back", TODAY)["intent"] == "back_received"
    back = answer_question(
        book,
        {"intent": "back_received", "date_filter": {"type": "all"}, "mode": "list"},
        TODAY,
    )["reply"]
    assert "Back 4" in back
    assert "1. RTO-2026-10-08 · 4" in back
    back_card = answer_question(
        book,
        {"intent": "back_received", "date_filter": {"type": "all"}, "mode": "list"},
        TODAY,
    )["table"]
    assert back_card["title"] == "Back"
    assert back_card["columns"] == ["Back"]


def test_delivered_dc_question_lists_each_number_and_the_sum() -> None:
    parsed = understand("ab tak delivered hue sare DC ka count", TODAY)
    assert parsed["intent"] == "delivered_dc"
    assert parsed["terms"] == []
    assert parsed["date_filter"]["type"] == "all"
    assert parsed["mode"] == "list"

    book = _book()
    reply = answer_question(
        book,
        {"intent": "delivered_dc", "date_filter": {"type": "all"}, "mode": "list"},
        TODAY,
    )["reply"]
    assert "DC count 2" in reply
    assert "Product wise sum 38 qty" in reply
    assert "1. Delivered-2026-09-23 · 21" in reply
    assert "2. Delivered-2026-10-01 · 17" in reply
    assert reply.count("\n") >= reply.split("DC count 2")[1].count("1. ")
    card = answer_question(
        book,
        {"intent": "delivered_dc", "date_filter": {"type": "all"}, "mode": "list"},
        TODAY,
    )["table"]
    assert card["title"] == "DC In"
    assert card["columns"] == ["DC In"]


def test_delivered_dc_with_store_does_not_offer_a_store_list(monkeypatch) -> None:
    _patch(monkeypatch)

    missing = ask_reconciliation("ab tak delivered hue sare DC ka count", today=TODAY)
    answer = ask_reconciliation(
        "lavkush mustard ab tak delivered hue sare DC ka count",
        today=TODAY,
    )

    assert missing["handled"] is True
    assert "Kaun sa warehouse? Naam Select Karo." in missing["reply"]
    assert "Lavkush Dark Store Barabanki (LAVK)" in missing["suggestions"]
    assert "Fatehpur Dark Store (FTP)" in missing["suggestions"]
    assert "DC count 2" in answer["reply"]
    assert "Product wise sum 38 qty" in answer["reply"]
    assert "1. OUT/1 · 21" in answer["reply"]
    assert "2. OUT/2 · 17" in answer["reply"]
    assert answer["context"]["candidates"] == []


def test_unknown_location_asks_to_select_a_warehouse_name(monkeypatch) -> None:
    _patch(monkeypatch)

    answer = ask_reconciliation("xyzpur mein sale out kitna hai", today=TODAY)
    selected = ask_reconciliation(answer["suggestions"][0], answer["context"], today=TODAY)

    assert answer["reply"].startswith("Kaun sa warehouse? Naam Select Karo.")
    assert "On hand" not in answer["reply"]
    assert answer["context"]["pending"] == "warehouse"
    assert selected["context"]["warehouses"][0]["id"] == 8
    assert "Kaun sa warehouse?" not in selected["reply"]


def test_delivered_dc_works_for_any_store_and_product(monkeypatch) -> None:
    _patch(monkeypatch)

    answer = ask_reconciliation(
        "fatehpur gramino ab tak delivered hue sare DC ka count",
        today=TODAY,
    )

    assert "Fatehpur Dark Store" in answer["reply"]
    assert "Gramino" in answer["reply"]
    assert "Lavkush" not in answer["reply"]
    assert "DC count 2" in answer["reply"]
    assert "Product wise sum 38 qty" in answer["reply"]
    assert answer["context"]["warehouses"][0]["id"] == 8
    assert answer["context"]["product"]["id"] == 9


def _warehouse(warehouse_id: int, name: str, code: str, stock_id: int) -> dict:
    return {
        "id": warehouse_id,
        "name": name,
        "code": code,
        "lot_stock_id": [stock_id, f"{code}/Stock"],
        "active": True,
    }


def _fake_search(model: str, domain: list, fields: list, order: str = "id asc", limit: int | None = None):
    blob = str(domain)
    if model == "stock.warehouse" and "ilike" in blob:
        if "fatehpur" in blob:
            return [_warehouse(8, "Fatehpur Dark Store", "FTP", 12)]
        if "lavkush" in blob or "barabanki" in blob:
            return [_warehouse(7, "Lavkush Dark Store Barabanki", "LAVK", 11)]
        return []
    if model == "stock.warehouse":
        wanted = next(
            (item[2] for item in domain if isinstance(item, list) and item[0] == "id"),
            None,
        )
        rows = [
            _warehouse(8, "Fatehpur Dark Store", "FTP", 12),
            _warehouse(7, "Lavkush Dark Store Barabanki", "LAVK", 11),
        ]
        if wanted is None:
            return rows
        return [row for row in rows if row["id"] == wanted]
    if model == "stock.quant" and "12" in blob:
        return [{"quantity": 9, "reserved_quantity": 1}]
    if model == "stock.quant":
        return [{"quantity": 2, "reserved_quantity": 2}]
    return []


def _fake_moves(_warehouse, _product_id):
    return [
        MoveRow("2026-09-19", "IN/1", "IN", "done", "", "Hub", 36, 36),
        MoveRow("2026-09-23", "OUT/1", "Delivered", "done", "AGRI1", "Ram", 21, 21),
        MoveRow("2026-10-01", "OUT/2", "Delivered", "done", "AGRI2", "Shiv", 17, 17),
        MoveRow("2026-10-08", "RTO/1", "RTO", "done", "AGRI1", "Ram", 4, 4),
    ]


def _fake_products(words: list[str]):
    if "gramino" in words:
        return [ProductRow(9, "Gramino", "200", "Gramino · 200")]
    if "musturd" in words or "1445" in words:
        return [ProductRow(101, "Musturd 1666", "1445", "Musturd 1666 · 1445")]
    return []


def _patch(monkeypatch) -> None:
    monkeypatch.setattr(odoo_ledger, "search_read_all", _fake_search)
    monkeypatch.setattr(odoo_ledger, "load_moves", _fake_moves)
    monkeypatch.setattr(odoo_ledger, "search_products", _fake_products)


def test_live_stock_reply_is_only_the_stock_card() -> None:
    from datetime import datetime
    from zoneinfo import ZoneInfo

    from app.modules.webtechAIAgent.reconciliation.answer_formatting import stock_card

    reply = stock_card(
        location="Barabanki (Lavkhush) - Dark Warehouse",
        product_name="Musturd 1666",
        sku="1445",
        variant="500 Gm",
        updated_at=datetime(2026, 10, 14, 18, 30, tzinfo=ZoneInfo("UTC")),
        snapshot=TODAY,
        on_hand=2,
        reserved=2,
        incoming=0,
    )
    assert reply == "\n".join(
        [
            "Location",
            "Barabanki (Lavkhush) - Dark Warehouse",
            "",
            "Product",
            "Musturd 1666 (SKU 1445)",
            "",
            "Variant",
            "500 Gm",
            "",
            "Stock Updated At",
            "15 Oct 2026, 12:00 AM",
            "",
            "On Hand",
            "2",
            "",
            "Reserved",
            "2",
            "",
            "Free",
            "0",
            "",
            "Incoming",
            "0",
        ]
    )
    lowered = reply.lower()
    assert "source:" not in lowered
    assert "live stock" not in lowered
    assert "free = on hand" not in lowered
    assert reply.count("On Hand") == 1


def test_any_warehouse_and_product_come_from_odoo(monkeypatch) -> None:
    _patch(monkeypatch)

    fatehpur = ask_reconciliation("Fatehpur mein Gramino kitna hai", today=TODAY)
    lavkush = ask_reconciliation("lavkush mustard on hand", today=TODAY)

    assert "workbook" not in fatehpur["reply"].lower()
    assert "Fatehpur Dark Store" in fatehpur["reply"]
    assert "Gramino" in fatehpur["reply"]
    assert "On Hand\n9" in fatehpur["reply"]
    assert "Source:" not in fatehpur["reply"]
    assert "Lavkush Dark Store" in lavkush["reply"]
    assert "On Hand\n2" in lavkush["reply"]
    assert "free = on hand minus reserved" not in lavkush["reply"]
    assert fatehpur["context"]["warehouses"][0]["id"] == 8
    assert lavkush["context"]["product"]["id"] == 101


def test_two_warehouses_are_both_answered(monkeypatch) -> None:
    _patch(monkeypatch)

    answer = ask_reconciliation("fatehpur aur lavkush ka mustard stock batao", today=TODAY)

    assert "Fatehpur Dark Store" in answer["reply"]
    assert "Lavkush Dark Store" in answer["reply"]
    assert "On Hand\n9" in answer["reply"]
    assert "On Hand\n2" in answer["reply"]


def test_store_and_city_together_pick_the_named_store(monkeypatch) -> None:
    def fake_search(model: str, domain: list, fields: list, order: str = "id asc", limit: int | None = None):
        blob = str(domain)
        if model == "stock.warehouse" and "ilike" in blob:
            if "lavkush" in blob or "lavkhush" in blob:
                return [_warehouse(27, "Barabanki (Lavkhush) - Dark Warehouse", "BL-D", 11)]
            if "barabanki" in blob:
                return [
                    _warehouse(21, "Barabanki (Indresh) - Dark Warehouse", "BI-D", 21),
                    _warehouse(27, "Barabanki (Lavkhush) - Dark Warehouse", "BL-D", 11),
                    _warehouse(37, "DS#37 - Barabanki (Fatehpur) - Gramik", "BB-DS", 37),
                    _warehouse(38, "DS#38 - Barabanki (Ramnagar) - Gramik", "BR-DS", 38),
                ]
            return []
        if model == "stock.warehouse":
            wanted = next(item[2] for item in domain if isinstance(item, list) and item[0] == "id")
            return [_warehouse(wanted, "Barabanki (Lavkhush) - Dark Warehouse", "BL-D", 11)]
        if model == "stock.quant":
            return [{"quantity": 2, "reserved_quantity": 2}]
        return []

    monkeypatch.setattr(odoo_ledger, "search_read_all", fake_search)
    monkeypatch.setattr(odoo_ledger, "load_moves", _fake_moves)
    monkeypatch.setattr(odoo_ledger, "search_products", _fake_products)

    answer = ask_reconciliation("Lavkush dark store barabanki me muster 1666 kitne hai", today=TODAY)

    assert "kai warehouse" not in answer["reply"]
    assert "Lavkhush" in answer["reply"]
    assert "On Hand\n2" in answer["reply"]
    assert answer["context"]["warehouses"][0]["id"] == 27


def test_follow_up_keeps_product_and_changes_the_metric(monkeypatch) -> None:
    _patch(monkeypatch)

    stock = ask_reconciliation("lavkush mustard kitna hai", today=TODAY)
    sales = ask_reconciliation("1 oct 2026 ko kitna sale hua", stock["context"], today=TODAY)
    other_store = ask_reconciliation("aur fatehpur", stock["context"], today=TODAY)

    assert "Total 17 qty" in sales["reply"]
    assert sales["context"]["product"]["sku"] == "1445"
    assert "On Hand\n9" in other_store["reply"]
    assert other_store["context"]["product"]["id"] == 101


def test_follow_up_without_a_new_name_keeps_product_and_warehouse(monkeypatch) -> None:
    _patch(monkeypatch)

    stock = ask_reconciliation("lavkush mustard kitna hai", today=TODAY)
    inward = ask_reconciliation(
        "Abhi tak total transfer in qty kitne hai lavkush pe",
        stock["context"],
        today=TODAY,
    )

    assert "nahi mila" not in inward["reply"]
    assert "Total 36 qty" in inward["reply"]
    assert inward["context"]["product"]["sku"] == "1445"
    assert inward["context"]["warehouses"][0]["id"] == 7


def test_misspelled_product_offers_close_names(monkeypatch) -> None:
    def fake_search(model: str, domain: list, fields: list, order: str = "id asc", limit: int | None = None):
        blob = str(domain)
        if model == "stock.warehouse" and ("lavkush" in blob or "lavkhush" in blob):
            return [_warehouse(7, "Lavkush Dark Store Barabanki", "LAVK", 11)]
        if model == "product.product":
            return [
                {
                    "id": 1445,
                    "name": "Musturd 1666",
                    "default_code": "1445",
                    "display_name": "[1445] Musturd 1666 (500 Gm)",
                },
                {
                    "id": 282,
                    "name": "Musturd 1666",
                    "default_code": "282",
                    "display_name": "[282] Musturd 1666 (1 Kg)",
                },
                {
                    "id": 185,
                    "name": "414 Mustard Seed",
                    "default_code": "185",
                    "display_name": "[185] 414 Mustard Seed",
                },
            ]
        if model == "stock.quant":
            return [{"quantity": 2, "reserved_quantity": 0}]
        return []

    monkeypatch.setattr(odoo_ledger, "search_read_all", fake_search)
    monkeypatch.setattr(odoo_ledger, "search_products", lambda words: [])
    answer = ask_reconciliation("lavkush mein mustered 1666 kitna hai", today=TODAY)

    assert "exact naam nahi mila" in answer["reply"]
    assert "Odoo" not in answer["reply"]
    assert "Musturd 1666 (500 Gm)" not in answer["reply"]
    assert "Musturd 1666 (1 Kg)" not in answer["reply"]
    assert "Number ya naam likho" not in answer["reply"]
    assert "Mustard Seed" not in answer["reply"]
    assert answer["suggestions"] == [
        "[1445] Musturd 1666 (500 Gm) · 1445",
        "[282] Musturd 1666 (1 Kg) · 282",
    ]
    assert answer["context"]["pending"] == "product"
    assert answer["context"]["candidates"][0]["sku"] == "1445"


def test_unknown_name_is_not_replaced_with_another_product(monkeypatch) -> None:
    _patch(monkeypatch)

    answer = ask_reconciliation("lavkush mein zinc kitna hai", today=TODAY)

    assert "nahi mila" in answer["reply"]
    assert "Odoo" not in answer["reply"]
    assert "On hand" not in answer["reply"]
    assert answer["context"].get("product") is None


def test_unknown_product_does_not_list_other_store_products(monkeypatch) -> None:
    def fake_search(model: str, domain: list, fields: list, order: str = "id asc", limit: int | None = None):
        blob = str(domain)
        if model == "stock.warehouse" and "ilike" in blob:
            if "lavkush" in blob:
                return [_warehouse(7, "Lavkush Dark Store Barabanki", "LAVK", 11)]
            return []
        if model == "stock.warehouse":
            return [_warehouse(7, "Lavkush Dark Store Barabanki", "LAVK", 11)]
        if model == "stock.quant":
            return [
                {"product_id": [101, "Musturd 1666"], "quantity": 4},
                {"product_id": [202, "Gramino"], "quantity": 1},
            ]
        if model == "product.product" and "ilike" in blob:
            return []
        if model == "product.product":
            return [
                {
                    "id": 101,
                    "name": "Musturd 1666",
                    "default_code": "1445",
                    "display_name": "[1445] Musturd 1666 (500 Gm)",
                },
                {
                    "id": 202,
                    "name": "Gramino",
                    "default_code": "200",
                    "display_name": "[200] Gramino",
                },
            ]
        return []

    monkeypatch.setattr(odoo_ledger, "search_read_all", fake_search)
    monkeypatch.setattr(odoo_ledger, "search_products", lambda words: [])
    answer = ask_reconciliation("lavkush mein zinc kitna hai", today=TODAY)

    assert "Product nahi mila: zinc" in answer["reply"]
    assert "Musturd" not in answer["reply"]
    assert "Gramino" not in answer["reply"]
    assert "Number ya naam likho" not in answer["reply"]
    assert "On hand" not in answer["reply"]
    assert answer["context"].get("pending") in {None, ""}
    assert answer["context"].get("product") is None
    assert answer["suggestions"] == []


def test_lmd_is_not_taken_from_sales() -> None:
    answer = ask_reconciliation("aaj LMD ne kitna pick kiya", today=TODAY)

    assert "LMD" in answer["reply"]
    assert "38" not in answer["reply"]
    assert "Total" not in answer["reply"]


def test_code_question_is_not_handled() -> None:
    answer = ask_reconciliation("rename the title on the login screen", today=TODAY)
    assert answer["handled"] is False


def test_all_products_excel_uses_the_same_tabs(monkeypatch, tmp_path) -> None:
    def fake_search(model: str, domain: list, fields: list, order: str = "id asc"):
        if model == "stock.move":
            return [
                {
                    "id": 1,
                    "product_id": [101, "Musturd 1666"],
                    "reference": "BL-D/IN/1",
                    "origin": "IBT",
                    "state": "done",
                    "date": "2026-09-19 05:00:00",
                    "product_uom_qty": 10,
                    "quantity": 10,
                    "location_id": [99, "Hub"],
                    "location_dest_id": [11, "BL-D/Stock"],
                    "picking_id": [5, "BL-D/IN/1"],
                    "sale_line_id": False,
                    "origin_returned_move_id": False,
                },
                {
                    "id": 2,
                    "product_id": [101, "Musturd 1666"],
                    "reference": "BL-D/OUT/1",
                    "origin": "AGRI1",
                    "state": "done",
                    "date": "2026-10-01 05:00:00",
                    "product_uom_qty": 4,
                    "quantity": 4,
                    "location_id": [11, "BL-D/Stock"],
                    "location_dest_id": [50, "Customers"],
                    "picking_id": [6, "BL-D/OUT/1"],
                    "sale_line_id": False,
                    "origin_returned_move_id": False,
                },
            ]
        if model == "stock.picking":
            return [
                {
                    "id": 5,
                    "name": "BL-D/IN/1",
                    "origin": "IBT",
                    "state": "done",
                    "partner_id": [1, "Hub"],
                    "scheduled_date": "2026-09-19 05:00:00",
                    "date_done": "2026-09-19 05:30:00",
                    "return_id": False,
                },
                {
                    "id": 6,
                    "name": "BL-D/OUT/1",
                    "origin": "AGRI1",
                    "state": "done",
                    "partner_id": [2, "Ram"],
                    "scheduled_date": "2026-10-01 05:00:00",
                    "date_done": "2026-10-01 05:30:00",
                    "return_id": False,
                },
            ]
        if model == "stock.location":
            return [
                {"id": 11, "usage": "internal"},
                {"id": 99, "usage": "supplier"},
                {"id": 50, "usage": "customer"},
            ]
        if model == "stock.quant":
            return [
                {
                    "id": 7,
                    "product_id": [101, "Musturd 1666"],
                    "location_id": [11, "BL-D/Stock"],
                    "lot_id": [3, "LOT1"],
                    "quantity": 6,
                    "reserved_quantity": 0,
                    "in_date": "2026-09-19 05:30:00",
                    "company_id": [1, "Uttar Pradesh"],
                }
            ]
        if model == "product.product":
            return [{"id": 101, "default_code": "1445", "name": "Musturd 1666"}]
        return []

    monkeypatch.setattr(excel_export, "search_read_all", fake_search)
    warehouse = Warehouse(7, "Barabanki (Lavkhush) - Dark Warehouse", "BL-D", 11, "BL-D/Stock", 12)
    target = tmp_path / "BL-D_all_products_reconciliation.xlsx"
    result = excel_export.export_store_workbook(warehouse, target)

    from openpyxl import load_workbook

    book = load_workbook(target)
    assert book.sheetnames == [
        "Summary", "Ledger", "IN", "Delivered", "RTO", "In Transit", "Outgoing", "Cancelled", "Live Stock",
    ]
    assert book["Summary"]["D6"].value.startswith("=SUMIFS")
    assert book["Summary"]["N6"].value == "=D6-E6+F6+J6-K6+L6-M6"
    assert book["Ledger"]["D2"].value == "BL-D/IN/1"
    assert book["IN"]["D2"].value == "BL-D/IN/1"
    assert book["Delivered"]["J2"].value == 4
    assert result["gaps"] == 0
    assert result["products"] == 1
    book.close()


def test_chat_asks_for_a_store_excel(monkeypatch, tmp_path) -> None:
    _patch(monkeypatch)
    sent: dict[str, object] = {}

    def fake_export(warehouse_id: int, dest=None):
        return {
            "path": str(tmp_path / "BL-D_all_products_reconciliation.xlsx"),
            "products": 3,
            "moves": 12,
            "gaps": 1,
            "matched": 2,
        }

    def fake_send(recipients: list[str], paths: list[str]) -> str:
        sent["recipients"] = recipients
        sent["paths"] = paths
        return f"Excel BL-D_all_products_reconciliation.xlsx bhej diya: {', '.join(recipients)}."

    monkeypatch.setattr(
        "app.modules.webtechAIAgent.reconciliation.service.export_store_by_id",
        fake_export,
    )
    monkeypatch.setattr(
        "app.modules.webtechAIAgent.reconciliation.service.send_excel_files",
        fake_send,
    )
    answer = ask_reconciliation("lavkush ke saare products ka reconciliation excel banao", today=TODAY)

    assert answer["handled"] is True
    assert "BL-D_all_products_reconciliation.xlsx" in answer["reply"]
    assert "Summary, Ledger, IN, Delivered, RTO, In Transit, Outgoing, Cancelled, Live Stock" in answer["reply"]
    assert answer["context"]["warehouses"][0]["id"] == 7
    assert sent == {}

    mailed = ask_reconciliation(
        "lavkush ka puri report bana ke ravi@gramik.in kr do",
        today=TODAY,
    )
    assert sent["recipients"] == ["ravi@gramik.in"]
    assert "ravi@gramik.in" in mailed["reply"]
    assert mailed["context"]["warehouses"][0]["id"] == 7


def test_follow_up_mails_the_saved_excel(monkeypatch, tmp_path) -> None:
    sent: dict[str, object] = {}

    def fake_send(recipients: list[str], paths: list[str]) -> str:
        sent["recipients"] = recipients
        sent["paths"] = paths
        return "Excel BL-D_all_products_reconciliation.xlsx bhej diya: manoj@gramik.in."

    def fail_export(warehouse_id: int, dest=None):
        raise AssertionError("mail follow-up must not rebuild the workbook")

    monkeypatch.setattr(
        "app.modules.webtechAIAgent.reconciliation.service.send_excel_files",
        fake_send,
    )
    monkeypatch.setattr(
        "app.modules.webtechAIAgent.reconciliation.service.export_store_by_id",
        fail_export,
    )
    path = str(tmp_path / "BL-D_all_products_reconciliation.xlsx")
    answer = ask_reconciliation(
        "isko manoj@gramik.in par mail kar do",
        {
            "intent": "store_excel",
            "excel_files": [path],
            "warehouses": [{"id": 7, "name": "Lavkush", "code": "BL-D"}],
        },
        today=TODAY,
    )

    assert sent["recipients"] == ["manoj@gramik.in"]
    assert sent["paths"] == [path]
    assert "bhej diya" in answer["reply"]


def test_ses_mail_uses_the_fixed_subject_and_attaches_the_file(monkeypatch, tmp_path) -> None:
    from openpyxl import Workbook

    from app.modules.webtechAIAgent.reconciliation import excel_mail

    workbook = Workbook()
    workbook.active.title = "Summary"
    target = tmp_path / "BL-D_all_products_reconciliation.xlsx"
    workbook.save(target)
    workbook.close()
    captured: list[object] = []

    class FakeSMTP:
        def __init__(self, host: str, port: int, timeout: int = 30) -> None:
            assert host == "email-smtp.example.com"

        def __enter__(self):
            return self

        def __exit__(self, *args: object) -> bool:
            return False

        def starttls(self) -> None:
            return None

        def login(self, user: str, password: str) -> None:
            assert user == "ses-user"

        def send_message(self, message: object) -> None:
            captured.append(message)

    class Settings:
        ses_smtp_host = "email-smtp.example.com"
        ses_smtp_port = 587
        ses_smtp_user = "ses-user"
        ses_smtp_pass = "secret"
        ses_from_email = "noreply@gramik.in"
        ses_from_name = "WebTechAgent AI Assistant"

    monkeypatch.setattr(excel_mail, "OUT_DIR", tmp_path)
    monkeypatch.setattr(excel_mail, "get_settings", lambda: Settings())
    monkeypatch.setattr(excel_mail.smtplib, "SMTP", FakeSMTP)

    reply = excel_mail.send_excel_files(["manoj@gramik.in"], [str(target)])
    message = captured[0]

    assert "bhej diya" in reply
    assert message["Subject"] == "Excel Report Generated by WebTechAgent AI Assistant"
    body = message.get_body(preferencelist=("plain",)).get_content()
    assert "Please find attached the Excel file generated by the WebTechAgent AI Assistant" in body
    attachments = list(message.iter_attachments())
    assert attachments[0].get_filename() == "BL-D_all_products_reconciliation.xlsx"


def test_api_uses_the_odoo_resolver(monkeypatch) -> None:
    _patch(monkeypatch)

    response = client.post(
        "/api/v1/reconciliation/ask",
        json={"message": "barabanki mustard on hand", "context": None},
    )

    assert response.status_code == 200
    body = response.json()["data"]
    assert body["handled"] is True
    assert "On Hand\n2" in body["reply"]
    assert "Source:" not in body["reply"]
    assert body["source"]["system"] == "odoo"
