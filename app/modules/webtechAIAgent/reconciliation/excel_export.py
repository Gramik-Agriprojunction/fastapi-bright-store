"""All-products reconciliation workbook for any dark store.

The tabs match the single-product workbook: Summary, Ledger, IN, Delivered,
RTO, In Transit, Outgoing, Cancelled, and Live Stock. Each row carries the
product because one file holds the whole store.
"""

from __future__ import annotations

import re
from collections import defaultdict
from datetime import UTC, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.modules.webtechAIAgent.reconciliation.entity_normalization import norm
from app.modules.webtechAIAgent.reconciliation.odoo_ledger import _warehouse_record
from app.modules.webtechAIAgent.service import (
    Warehouse,
    classify,
    m2o_id,
    m2o_name,
    partner_name,
    qty_of,
    sale_ref,
    search_read_all,
)

OUT_DIR = Path("/Users/sourabhyadav/Documents/Project/Gramik-Live/data_reconciliation")
SHEETS = (
    "Summary",
    "Ledger",
    "IN",
    "Delivered",
    "RTO",
    "In Transit",
    "Outgoing",
    "Cancelled",
    "Live Stock",
)
STOCK_IN = {"IN", "RTO", "Adjustment IN", "Other IN"}
STOCK_OUT = {"Delivered", "Adjustment OUT", "Other OUT"}

NAVY = "1B3A4B"
TEAL = "0F6E56"
LINE = "D5DDE2"
FILLS = {
    "header": PatternFill("solid", fgColor=NAVY),
    "title": PatternFill("solid", fgColor=TEAL),
    "IN": PatternFill("solid", fgColor="E5F6EF"),
    "Delivered": PatternFill("solid", fgColor="FDECEC"),
    "RTO": PatternFill("solid", fgColor="FFF4E5"),
    "RTO in transit": PatternFill("solid", fgColor="F8E8D4"),
    "In transit": PatternFill("solid", fgColor="F3F0FA"),
    "Transit leg": PatternFill("solid", fgColor="F3F0FA"),
    "Outgoing": PatternFill("solid", fgColor="E8F1FB"),
    "Adjustment IN": PatternFill("solid", fgColor="E7F6EA"),
    "Adjustment OUT": PatternFill("solid", fgColor="FBEAEA"),
    "Other IN": PatternFill("solid", fgColor="EEF6F2"),
    "Other OUT": PatternFill("solid", fgColor="F8EEEE"),
    "Cancelled": PatternFill("solid", fgColor="F4F5F6"),
    "match": PatternFill("solid", fgColor="D8F3E4"),
    "gap": PatternFill("solid", fgColor="FDE2E1"),
    "live": PatternFill("solid", fgColor="E7F0F6"),
}
THIN = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
FONT_HEADER = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
FONT_TITLE = Font(name="Calibri", bold=True, color="FFFFFF", size=16)
FONT_LABEL = Font(name="Calibri", size=11, color="1A1A1A")
FONT_SMALL = Font(name="Calibri", size=9, color="5C6B73", italic=True)

LEDGER_HEADERS = [
    "Date (IST)", "SKU", "Product", "Document", "Bucket", "State", "Sale Order",
    "Customer / Partner", "From", "To", "Qty In", "Qty Out", "Running On Hand",
    "Return Of", "Lot", "Note", "Product Id",
]
DETAIL_HEADERS = [
    "Date (IST)", "SKU", "Product", "Document", "State", "Sale Order",
    "Customer / Partner", "From", "To", "Qty", "Demand", "Return Of", "Lot", "Origin",
]
SUMMARY_HEADERS = [
    "Product Id", "SKU", "Product", "IN", "Delivered", "RTO", "RTO in transit",
    "In transit", "Outgoing", "Adjustment IN", "Adjustment OUT", "Other IN", "Other OUT",
    "Hona chahiye", "Live on hand", "Reserved", "Free", "Gap", "Status",
]


def wants_store_excel(text: str) -> bool:
    hay = norm(text)
    full_report = any(
        phrase in hay
        for phrase in (
            "puri report",
            "poori report",
            "pura report",
            "poora report",
            "pure report",
            "full report",
        )
    )
    file_ask = any(word in hay for word in ("excel", "xlsx", "workbook"))
    all_products = any(
        phrase in hay
        for phrase in ("all product", "saare product", "sare product", "sab product", "sabhi product")
    )
    recon = "reconcil" in hay or "hisab" in hay
    return full_report or (file_ask and (all_products or recon))


def _parse_ist(value: object) -> datetime | None:
    if not value or value is False:
        return None
    text = str(value).split(".")[0].replace("T", " ")
    match = re.match(r"^(\d{4}-\d{2}-\d{2})[ T](\d{2}:\d{2}:\d{2})", text)
    if not match:
        return None
    utc = datetime.strptime(f"{match.group(1)} {match.group(2)}", "%Y-%m-%d %H:%M:%S").replace(tzinfo=UTC)
    return (utc + timedelta(minutes=330)).replace(tzinfo=None)


def _note(bucket: str) -> str:
    notes = {
        "IN": "Store stock me receive hua.",
        "Delivered": "Customer delivery done. Store stock se nikal gaya.",
        "RTO": "RTO receive ho gaya, stock wapas store me aaya.",
        "RTO in transit": "Return open hai, store me abhi receive nahi hua.",
        "In transit": "Abhi transit / receipt open hai. Store on-hand me nahi kata.",
        "Outgoing": "Delivery assign hai, done nahi. On-hand abhi ghat-ta nahi, reserved ho sakta hai.",
        "Adjustment IN": "Inventory adjustment se stock badha.",
        "Adjustment OUT": "Inventory adjustment se stock ghata.",
        "Other OUT": "Customer delivery ke alawa stock bahar gaya.",
        "Other IN": "Normal receipt ke alawa stock andar aaya.",
        "Cancelled": "Cancel. Stock move nahi hua. Demand alag column me hai.",
        "Transit leg": "Hub se transit. Store stock tab badhta hai jab IN receipt done ho.",
    }
    return notes.get(bucket, "")


def _paint(ws: Worksheet, header_row: int, cols: int, widths: dict[str, int]) -> None:
    for col in range(1, cols + 1):
        cell = ws.cell(header_row, col)
        cell.fill = FILLS["header"]
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
    headers = [ws.cell(header_row, col).value for col in range(1, cols + 1)]
    bucket_at = headers.index("Bucket") + 1 if "Bucket" in headers else None
    for row in range(header_row + 1, ws.max_row + 1):
        bucket = ws.cell(row, bucket_at).value if bucket_at else None
        fill = FILLS.get(str(bucket)) if bucket else None
        for col in range(1, cols + 1):
            cell = ws.cell(row, col)
            cell.border = THIN
            cell.font = FONT_LABEL
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if fill:
                cell.fill = fill
            if isinstance(cell.value, datetime):
                cell.number_format = "DD-MMM-YYYY HH:MM"
            if isinstance(cell.value, float) or (isinstance(cell.value, str) and cell.value.startswith("=")):
                cell.number_format = "0.##"
    for index, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(index)].width = widths.get(str(header), 16)
    last = get_column_letter(cols)
    ws.auto_filter.ref = f"A{header_row}:{last}{max(ws.max_row, header_row)}"
    ws.freeze_panes = f"A{header_row + 1}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"


def _chunks(ids: list[int], size: int = 80) -> list[list[int]]:
    return [ids[start : start + size] for start in range(0, len(ids), size)]


def _load(warehouse: Warehouse) -> tuple[list[dict[str, object]], list[list[object]], dict[int, float], dict[int, float], dict[int, float]]:
    location_ids = [item for item in (warehouse.stock_id, warehouse.transit_id) if item]
    moves = search_read_all(
        "stock.move",
        ["|", ["location_id", "in", location_ids], ["location_dest_id", "in", location_ids]],
        [
            "id", "product_id", "reference", "origin", "state", "date", "product_uom_qty",
            "quantity", "location_id", "location_dest_id", "picking_id", "sale_line_id",
            "origin_returned_move_id",
        ],
    )
    picking_ids = sorted({pid for move in moves if (pid := m2o_id(move.get("picking_id")))})
    pickings: dict[int, dict[str, object]] = {}
    for chunk in _chunks(picking_ids):
        for row in search_read_all(
            "stock.picking",
            [["id", "in", chunk]],
            ["id", "name", "origin", "state", "partner_id", "scheduled_date", "date_done", "return_id"],
        ):
            pickings[int(row["id"])] = row
    loc_ids = set(location_ids)
    for move in moves:
        for key in ("location_id", "location_dest_id"):
            loc_id = m2o_id(move.get(key))
            if loc_id:
                loc_ids.add(loc_id)
    usage: dict[int, str] = {}
    for chunk in _chunks(sorted(loc_ids)):
        for row in search_read_all(
            "stock.location",
            [["id", "in", chunk]],
            ["id", "usage"],
        ):
            usage[int(row["id"])] = str(row.get("usage") or "")
    lots_by_move: dict[int, set[str]] = defaultdict(set)
    for line in search_read_all(
        "stock.move.line",
        ["|", ["location_id", "in", location_ids], ["location_dest_id", "in", location_ids]],
        ["move_id", "lot_id"],
    ):
        move_id = m2o_id(line.get("move_id"))
        lot = m2o_name(line.get("lot_id"))
        if move_id and lot:
            lots_by_move[move_id].add(lot)
    quants = search_read_all(
        "stock.quant",
        [["location_id", "in", location_ids]],
        ["id", "product_id", "location_id", "lot_id", "quantity", "reserved_quantity", "in_date", "company_id"],
    )
    product_ids = {pid for move in moves if (pid := m2o_id(move.get("product_id")))}
    product_ids |= {pid for quant in quants if (pid := m2o_id(quant.get("product_id")))}
    products: dict[int, dict[str, object]] = {}
    for chunk in _chunks(sorted(product_ids)):
        for row in search_read_all(
            "product.product",
            [["id", "in", chunk]],
            ["id", "default_code", "name"],
        ):
            products[int(row["id"])] = row

    records: list[dict[str, object]] = []
    for move in moves:
        product_id = m2o_id(move.get("product_id"))
        if not product_id:
            continue
        picking = pickings.get(m2o_id(move.get("picking_id")) or 0)
        bucket = classify(move, picking, usage, warehouse.stock_id, warehouse.transit_id)
        if not bucket:
            continue
        product = products.get(product_id, {})
        qty = qty_of(move)
        demand = float(move.get("product_uom_qty") or 0)
        when = _parse_ist((picking or {}).get("date_done") if move.get("state") == "done" else (picking or {}).get("scheduled_date"))
        when = when or _parse_ist(move.get("date"))
        qty_in = qty if bucket in STOCK_IN else None
        qty_out = qty if bucket in STOCK_OUT or bucket in {"Outgoing", "RTO in transit"} else None
        returned = ""
        if picking and m2o_name(picking.get("return_id")):
            returned = m2o_name(picking.get("return_id"))
        elif m2o_name(move.get("origin_returned_move_id")):
            returned = m2o_name(move.get("origin_returned_move_id"))
        records.append(
            {
                "when": when,
                "product_id": product_id,
                "sku": str(product.get("default_code") or ""),
                "product": str(product.get("name") or m2o_name(move.get("product_id"))),
                "doc": str((picking or {}).get("name") or move.get("reference") or f"MOVE/{move.get('id')}"),
                "bucket": bucket,
                "state": str(move.get("state") or ""),
                "sale": sale_ref(move, picking),
                "customer": partner_name((picking or {}).get("partner_id")),
                "src": m2o_name(move.get("location_id")),
                "dst": m2o_name(move.get("location_dest_id")),
                "qty_in": qty_in,
                "qty_out": qty_out,
                "qty": qty,
                "demand": demand,
                "return_of": returned,
                "lot": ", ".join(sorted(lots_by_move.get(int(move["id"]), []))),
                "note": _note(bucket),
                "origin": str((picking or {}).get("origin") or move.get("origin") or ""),
            }
        )

    running: dict[int, float] = defaultdict(float)
    for row in sorted(records, key=lambda item: (int(item["product_id"]), item["when"] is None, str(item["when"] or ""), str(item["doc"]))):
        if row["bucket"] in STOCK_IN or row["bucket"] in STOCK_OUT:
            running[int(row["product_id"])] += float(row["qty_in"] or 0) - float(row["qty_out"] or 0)
            row["running"] = running[int(row["product_id"])]
        else:
            row["running"] = None

    on_hand: dict[int, float] = defaultdict(float)
    reserved: dict[int, float] = defaultdict(float)
    transit: dict[int, float] = defaultdict(float)
    live_rows: list[list[object]] = []
    for quant in quants:
        product_id = m2o_id(quant.get("product_id"))
        if not product_id:
            continue
        qty = float(quant.get("quantity") or 0)
        hold = float(quant.get("reserved_quantity") or 0)
        if not qty and not hold:
            continue
        product = products.get(product_id, {})
        loc = m2o_id(quant.get("location_id"))
        if loc == warehouse.stock_id:
            on_hand[product_id] += qty
            reserved[product_id] += hold
        elif loc == warehouse.transit_id:
            transit[product_id] += qty
        live_rows.append(
            [
                str(product.get("default_code") or ""),
                str(product.get("name") or ""),
                m2o_name(quant.get("location_id")),
                m2o_name(quant.get("lot_id")) or "—",
                qty,
                hold,
                qty - hold,
                _parse_ist(quant.get("in_date")),
                quant.get("id"),
                m2o_name(quant.get("company_id")),
            ]
        )
    return records, live_rows, on_hand, reserved, transit


def _safe_name(warehouse: Warehouse) -> str:
    code = re.sub(r"[^A-Za-z0-9_-]+", "-", warehouse.code or "store").strip("-") or "store"
    return f"{code}_all_products_reconciliation.xlsx"


def export_store_workbook(warehouse: Warehouse, dest: Path | None = None) -> dict[str, object]:
    records, live_rows, on_hand, reserved, transit = _load(warehouse)
    product_ids = set(on_hand) | set(reserved) | set(transit) | {int(row["product_id"]) for row in records}
    product_ids = {
        pid
        for pid in product_ids
        if abs(on_hand[pid]) > 1e-6
        or abs(reserved[pid]) > 1e-6
        or abs(transit[pid]) > 1e-6
        or any(int(row["product_id"]) == pid for row in records)
    }
    names = {int(row["product_id"]): (str(row["sku"]), str(row["product"])) for row in records}
    ordered = sorted(product_ids, key=lambda pid: (names.get(pid, ("", ""))[0] or "zzzz", pid))

    book = Workbook()
    ledger = book.active
    ledger.title = "Ledger"
    ledger.append(LEDGER_HEADERS)
    for row in sorted(records, key=lambda item: (str(item["sku"]), item["when"] is None, str(item["when"] or ""), str(item["doc"]))):
        ledger.append(
            [
                row["when"], row["sku"], row["product"], row["doc"], row["bucket"], row["state"],
                row["sale"], row["customer"], row["src"], row["dst"], row["qty_in"], row["qty_out"],
                row["running"], row["return_of"], row["lot"], row["note"], row["product_id"],
            ]
        )
    _paint(ledger, 1, len(LEDGER_HEADERS), {
        "Date (IST)": 20, "SKU": 14, "Product": 36, "Document": 22, "Bucket": 16, "State": 12,
        "Sale Order": 22, "Customer / Partner": 24, "From": 36, "To": 28, "Qty In": 12, "Qty Out": 12,
        "Running On Hand": 18, "Return Of": 22, "Lot": 18, "Note": 46, "Product Id": 12,
    })

    summary = book.create_sheet("Summary", 0)
    summary.merge_cells("A1:S1")
    summary["A1"] = f"{warehouse.name}  —  all products stock reconciliation"
    summary["A1"].font = FONT_TITLE
    summary["A1"].fill = FILLS["title"]
    summary.row_dimensions[1].height = 28
    for col in range(1, 20):
        summary.cell(1, col).fill = FILLS["title"]
    summary.merge_cells("A2:S2")
    stamp = datetime.now(UTC) + timedelta(minutes=330)
    summary["A2"] = (
        f"{warehouse.stock_name or warehouse.code}  |  {warehouse.name}  |  "
        f"Live stock  |  read-only  |  {stamp.strftime('%d %b %Y %H:%M')} IST"
    )
    summary["A2"].font = FONT_SMALL
    summary.merge_cells("A3:S3")
    summary["A3"] = (
        "Hona chahiye = IN + Adjustment IN + RTO − Delivered − Adjustment OUT + Other IN − Other OUT. "
        "Outgoing aur cancelled demand stock se nahi kat-te. In transit location qty store on-hand me nahi jodi jaati. "
        "Gap = Live on hand − Hona chahiye."
    )
    summary["A3"].alignment = Alignment(wrap_text=True, vertical="center")
    summary.row_dimensions[3].height = 34
    for col, header in enumerate(SUMMARY_HEADERS, 1):
        summary.cell(5, col, header)
    first = 6
    for offset, product_id in enumerate(ordered):
        row = first + offset
        sku, name = names.get(product_id, ("", ""))
        summary.cell(row, 1, product_id)
        summary.cell(row, 2, sku)
        summary.cell(row, 3, name)
        summary.cell(row, 4, f'=SUMIFS(Ledger!K:K,Ledger!Q:Q,A{row},Ledger!E:E,"IN")')
        summary.cell(row, 5, f'=SUMIFS(Ledger!L:L,Ledger!Q:Q,A{row},Ledger!E:E,"Delivered")')
        summary.cell(row, 6, f'=SUMIFS(Ledger!K:K,Ledger!Q:Q,A{row},Ledger!E:E,"RTO")')
        summary.cell(row, 7, f'=SUMIFS(Ledger!L:L,Ledger!Q:Q,A{row},Ledger!E:E,"RTO in transit")')
        summary.cell(row, 8, round(transit[product_id], 4))
        summary.cell(row, 9, f'=SUMIFS(Ledger!L:L,Ledger!Q:Q,A{row},Ledger!E:E,"Outgoing")')
        summary.cell(row, 10, f'=SUMIFS(Ledger!K:K,Ledger!Q:Q,A{row},Ledger!E:E,"Adjustment IN")')
        summary.cell(row, 11, f'=SUMIFS(Ledger!L:L,Ledger!Q:Q,A{row},Ledger!E:E,"Adjustment OUT")')
        summary.cell(row, 12, f'=SUMIFS(Ledger!K:K,Ledger!Q:Q,A{row},Ledger!E:E,"Other IN")')
        summary.cell(row, 13, f'=SUMIFS(Ledger!L:L,Ledger!Q:Q,A{row},Ledger!E:E,"Other OUT")')
        summary.cell(row, 14, f"=D{row}-E{row}+F{row}+J{row}-K{row}+L{row}-M{row}")
        summary.cell(row, 15, round(on_hand[product_id], 4))
        summary.cell(row, 16, round(reserved[product_id], 4))
        summary.cell(row, 17, f"=O{row}-P{row}")
        summary.cell(row, 18, f"=O{row}-N{row}")
        summary.cell(row, 19, f'=IF(ABS(R{row})<0.001,"Matched",IF(R{row}<0,"Short","Excess"))')
    last = first + len(ordered) - 1 if ordered else first - 1
    if ordered:
        total = last + 1
        summary.cell(total, 3, "TOTAL")
        for col in range(4, 19):
            letter = get_column_letter(col)
            summary.cell(total, col, f"=SUM({letter}{first}:{letter}{last})")
        summary.cell(total, 19, f'=COUNTIF(S{first}:S{last},"Matched")&" matched"')
        for col in range(1, 20):
            cell = summary.cell(total, col)
            cell.font = Font(name="Calibri", bold=True, size=11)
            cell.fill = FILLS["live"]
            cell.border = THIN
    _paint(summary, 5, len(SUMMARY_HEADERS), {header: 14 for header in SUMMARY_HEADERS} | {"Product": 36, "Product Id": 12})
    if ordered:
        summary.auto_filter.ref = f"A5:S{last}"
        summary.freeze_panes = "D6"
        summary.conditional_formatting.add(
            f"S{first}:S{last}",
            FormulaRule(formula=[f'S{first}="Matched"'], fill=FILLS["match"]),
        )
        summary.conditional_formatting.add(
            f"R{first}:R{last}",
            CellIsRule(operator="greaterThan", formula=["0.001"], fill=FILLS["gap"]),
        )
        summary.conditional_formatting.add(
            f"R{first}:R{last}",
            CellIsRule(operator="lessThan", formula=["-0.001"], fill=FILLS["gap"]),
        )

    def add_sheet(title: str, buckets: set[str]) -> Worksheet:
        ws = book.create_sheet(title)
        ws.append(DETAIL_HEADERS)
        subset = [row for row in records if row["bucket"] in buckets]
        subset.sort(key=lambda item: (str(item["sku"]), item["when"] is None, str(item["when"] or ""), str(item["doc"])))
        for row in subset:
            qty = 0.0 if row["bucket"] == "Cancelled" else row["qty"]
            if row["bucket"] == "Transit leg":
                qty = row["qty"]
            ws.append(
                [
                    row["when"], row["sku"], row["product"], row["doc"], row["state"], row["sale"],
                    row["customer"], row["src"], row["dst"], qty, row["demand"], row["return_of"],
                    row["lot"], row["origin"],
                ]
            )
        if ws.max_row == 1:
            ws.append(["", "", "No rows"] + [""] * 11)
        _paint(ws, 1, len(DETAIL_HEADERS), {
            "Date (IST)": 20, "SKU": 14, "Product": 36, "Document": 22, "State": 12,
            "Sale Order": 22, "Customer / Partner": 24, "From": 36, "To": 28, "Qty": 12,
            "Demand": 12, "Return Of": 22, "Lot": 18, "Origin": 24,
        })
        return ws

    add_sheet("IN", {"IN"})
    add_sheet("Delivered", {"Delivered"})
    add_sheet("RTO", {"RTO"})
    transit_sheet = add_sheet("In Transit", {"In transit", "Transit leg"})
    if transit_sheet.max_row == 2 and transit_sheet.cell(2, 3).value == "No rows":
        transit_sheet.delete_rows(2)
    for product_id in ordered:
        if abs(transit[product_id]) < 1e-6:
            continue
        sku, name = names.get(product_id, ("", ""))
        transit_sheet.append(
            [
                None, sku, name, "TRANSIT LOCATION NOW", "on hand", "", "",
                warehouse.name, "", transit[product_id], transit[product_id], "", "",
                "Live transit location. Store on-hand me shamil nahi.",
            ]
        )
    if transit_sheet.max_row == 1:
        transit_sheet.append(["", "", "In transit quantity 0."] + [""] * 11)
    _paint(transit_sheet, 1, len(DETAIL_HEADERS), {"Product": 36, "Date (IST)": 20, "Document": 24})
    add_sheet("Outgoing", {"Outgoing"})
    add_sheet("Cancelled", {"Cancelled"})

    live = book.create_sheet("Live Stock")
    live_headers = ["SKU", "Product", "Location", "Lot", "On Hand", "Reserved", "Free", "In Date (IST)", "Quant ID", "Company"]
    live.append(live_headers)
    for row in sorted(live_rows, key=lambda item: (str(item[0]), str(item[2]), str(item[3]))):
        live.append(row)
    if live.max_row == 1:
        live.append(["", "No quant", "", "", 0, 0, 0, None, "", ""])
    _paint(live, 1, len(live_headers), {
        "SKU": 14, "Product": 36, "Location": 42, "Lot": 18, "On Hand": 12, "Reserved": 12,
        "Free": 12, "In Date (IST)": 20, "Quant ID": 12, "Company": 22,
    })

    target = dest or (OUT_DIR / _safe_name(warehouse))
    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(target)
    book.close()
    gaps = 0
    for product_id in ordered:
        inward = sum(float(row["qty_in"] or 0) for row in records if int(row["product_id"]) == product_id and row["bucket"] == "IN")
        sold = sum(float(row["qty_out"] or 0) for row in records if int(row["product_id"]) == product_id and row["bucket"] == "Delivered")
        returned = sum(float(row["qty_in"] or 0) for row in records if int(row["product_id"]) == product_id and row["bucket"] == "RTO")
        adj_in = sum(float(row["qty_in"] or 0) for row in records if int(row["product_id"]) == product_id and row["bucket"] == "Adjustment IN")
        adj_out = sum(float(row["qty_out"] or 0) for row in records if int(row["product_id"]) == product_id and row["bucket"] == "Adjustment OUT")
        other_in = sum(float(row["qty_in"] or 0) for row in records if int(row["product_id"]) == product_id and row["bucket"] == "Other IN")
        other_out = sum(float(row["qty_out"] or 0) for row in records if int(row["product_id"]) == product_id and row["bucket"] == "Other OUT")
        should = inward + adj_in + returned + other_in - sold - adj_out - other_out
        if abs(on_hand[product_id] - should) >= 0.001:
            gaps += 1
    return {
        "path": str(target),
        "products": len(ordered),
        "moves": len(records),
        "gaps": gaps,
        "matched": len(ordered) - gaps,
    }


def export_store_by_id(warehouse_id: int, dest: Path | None = None) -> dict[str, object] | str:
    warehouse = _warehouse_record(warehouse_id)
    if warehouse is None:
        return "Warehouse nahi mila."
    return export_store_workbook(warehouse, dest)
