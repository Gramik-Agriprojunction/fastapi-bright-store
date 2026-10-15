"""Excel workbooks for one inventory question. Totals are the filtered rows."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.modules.webtechAIAgent.reconciliation.excel_export import (
    FILLS,
    FONT_HEADER,
    FONT_LABEL,
    FONT_TITLE,
    OUT_DIR,
    THIN,
    _paint,
)
from app.modules.webtechAIAgent.reconciliation.intent_mapping import REPORT_INTENTS
from app.modules.webtechAIAgent.reconciliation.query_execution import execute_report
from app.modules.webtechAIAgent.reconciliation.workbook import Book

LINE = "D5DDE2"


def save_report(
    books: list[Book],
    context: dict[str, object],
    today: date,
    dest: Path | None = None,
) -> dict[str, object] | str:
    if not books:
        return "Excel ke liye warehouse data nahi mila."
    tables: list[dict[str, object]] = []
    record_count = 0
    intents = _intents(context)
    for book in books:
        for intent in intents:
            scoped = dict(context)
            scoped["intent"] = intent
            if context.get("complete"):
                scoped["columns"] = []
            result = execute_report(book, scoped, today)
            record_count += int(result.get("record_count") or 0)
            for table in result.get("tables") or []:
                if isinstance(table, dict):
                    tables.append(table)
    if not tables:
        return "Excel ke liye koi sheet nahi bani."
    merged = _merge(tables)
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    for table in merged:
        _write_sheet(workbook, table)
    stamp = datetime.now().strftime("%Y%m%d%H%M%S")
    warehouse = _safe(books[0].warehouse or books[0].store_name or "store")
    sku = _safe(books[0].sku or "all")
    intent = _safe(str(context.get("intent") or "report"))
    target = dest or (OUT_DIR / f"{warehouse}_{sku}_{intent}_{stamp}.xlsx")
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    workbook.close()
    return {
        "path": str(target),
        "name": target.name,
        "records": record_count,
        "sheets": [str(table["sheet"]) for table in merged],
    }


def validate_for_email(paths: list[str], record_count: int) -> str | None:
    if record_count <= 0:
        return "Workbook mein maange hue records nahi hain."
    for item in paths:
        path = Path(item)
        if not path.is_file():
            return "Excel file disk par nahi mili."
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except OSError as exc:
            return f"Excel file khuli nahi: {exc}"
        try:
            if not workbook.sheetnames:
                return "Excel file mein koi sheet nahi hai."
            headers = _header_names(workbook)
        finally:
            workbook.close()
        if not headers:
            return "Excel file mein header row nahi hai."
    return None


def _intents(context: dict[str, object]) -> list[str]:
    if context.get("complete"):
        return [
            "SALES_LEDGER",
            "DELIVERED_ORDERS",
            "RTO_RECEIVED_REPORT",
            "INCOMING_STOCK_REPORT",
            "OUTGOING_STOCK_REPORT",
            "IN_TRANSIT_REPORT",
            "CANCELLED_ORDERS_REPORT",
            "LIVE_STOCK_REPORT",
            "PRODUCT_STOCK_SUMMARY",
        ]
    intent = str(context.get("intent") or "")
    if intent in REPORT_INTENTS:
        return [intent]
    return ["SALES_LEDGER"]


def _header_names(workbook: object) -> list[str]:
    found: list[str] = []
    for name in workbook.sheetnames:
        sheet = workbook[name]
        first = next(sheet.iter_rows(max_row=1, values_only=True), None)
        if not first:
            continue
        found.extend(str(value) for value in first if value)
    return found


def _merge(tables: list[dict[str, object]]) -> list[dict[str, object]]:
    order: list[str] = []
    grouped: dict[str, dict[str, object]] = {}
    for table in tables:
        name = str(table.get("sheet") or "Report")
        if name == "Summary" and name in grouped:
            continue
        if name not in grouped:
            order.append(name)
            grouped[name] = {
                "sheet": name,
                "columns": list(table.get("columns") or []),
                "headers": list(table.get("headers") or []),
                "records": [],
            }
        current = grouped[name]
        seen = {
            (
                record.get("order_number"),
                record.get("document_number"),
                record.get("sku"),
                record.get("sale_out"),
                record.get("rto_in"),
                record.get("delivered_qty"),
                record.get("received_qty"),
                record.get("transaction_state"),
            )
            for record in current["records"]
            if isinstance(record, dict)
        }
        for record in table.get("records") or []:
            if not isinstance(record, dict):
                continue
            key = (
                record.get("order_number"),
                record.get("document_number"),
                record.get("sku"),
                record.get("sale_out"),
                record.get("rto_in"),
                record.get("delivered_qty"),
                record.get("received_qty"),
                record.get("transaction_state"),
            )
            if key in seen:
                continue
            seen.add(key)
            current["records"].append(record)
    return [grouped[name] for name in order]


def _write_sheet(workbook: Workbook, table: dict[str, object]) -> None:
    title = str(table["sheet"])[:31]
    sheet = workbook.create_sheet(title)
    headers = [str(item) for item in table.get("headers") or []]
    columns = [str(item) for item in table.get("columns") or []]
    records = [row for row in table.get("records") or [] if isinstance(row, dict)]
    if title == "Summary":
        _write_summary(sheet, headers, columns, records)
        return
    sheet.append(headers)
    numeric = {
        "Sale Out", "RTO IN", "Delivered Qty", "Received Qty", "Incoming Transfer",
        "Outgoing Transfer", "On Hand Qty", "Reserved Qty", "Free Qty", "Expected Qty",
        "Variance", "Cancelled Qty",
    }
    for record in records:
        sheet.append([_cell(record.get(column)) for column in columns])
    if records and any(header in numeric for header in headers):
        total: list[object] = []
        for index, header in enumerate(headers):
            if index == 0:
                total.append("TOTAL")
                continue
            if header not in numeric:
                total.append(None)
                continue
            letter = get_column_letter(index + 1)
            total.append(f"=SUM({letter}2:{letter}{len(records) + 1})")
        sheet.append(total)
    widths = {header: 18 for header in headers}
    widths.update({"Product Name": 28, "Document Number": 28, "Warehouse": 28, "Destination": 24})
    if headers:
        _paint(sheet, 1, len(headers), widths)


def _write_summary(
    sheet: Worksheet,
    headers: list[str],
    columns: list[str],
    records: list[dict[str, object]],
) -> None:
    sheet["A1"] = "Inventory report"
    sheet["A1"].font = FONT_TITLE
    sheet["A1"].fill = PatternFill("solid", fgColor="0F6E56")
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, len(headers) or 2))
    orders = {
        str(record.get("order_number") or record.get("document_number") or "")
        for record in records
        if record.get("order_number") or record.get("document_number")
    }
    sheet["A3"] = "Orders"
    sheet["B3"] = len(orders) if orders else len(records)
    row_at = 4
    numeric_headers = {
        "Sale Out", "RTO IN", "Delivered Qty", "Received Qty", "Incoming Transfer",
        "Outgoing Transfer", "On Hand Qty", "Reserved Qty", "Free Qty", "Expected Qty",
        "Variance", "Cancelled Qty",
    }
    for header, column in zip(headers, columns):
        if header not in numeric_headers:
            continue
        amount = 0.0
        for record in records:
            value = record.get(column)
            if isinstance(value, (int, float)):
                amount += float(value)
        sheet.cell(row_at, 1, header)
        sheet.cell(row_at, 2, amount)
        sheet.cell(row_at, 2).number_format = "0.##"
        row_at += 1
    thin = Border(
        left=Side(style="thin", color=LINE),
        right=Side(style="thin", color=LINE),
        top=Side(style="thin", color=LINE),
        bottom=Side(style="thin", color=LINE),
    )
    for row in sheet.iter_rows(min_row=3, max_row=max(3, row_at - 1), max_col=2):
        for cell in row:
            cell.border = thin
            cell.font = FONT_LABEL
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 16
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A3:B{max(3, row_at - 1)}"
    sheet.sheet_view.showGridLines = False
    header = sheet["A3"]
    header.fill = FILLS["header"]
    header.font = FONT_HEADER
    sheet["B3"].fill = FILLS["header"]
    sheet["B3"].font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")


def _cell(value: object) -> object:
    if value is None or value == "":
        return None
    return value


def _safe(value: str) -> str:
    cleaned = "".join(char if char.isalnum() or char in "-_" else "_" for char in value)
    return cleaned.strip("_") or "report"


__all__ = ["OUT_DIR", "save_report", "validate_for_email"]
