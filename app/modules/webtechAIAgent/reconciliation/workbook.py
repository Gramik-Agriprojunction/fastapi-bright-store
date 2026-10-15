"""Read the reconciliation workbook. Never writes it."""

from __future__ import annotations

import re
import warnings
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from app.core.config import get_settings

SNAPSHOT_RE = re.compile(r"(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})")
MONTHS = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}


@dataclass
class Move:
    when: datetime
    doc: str
    bucket: str
    state: str
    sale: str
    partner: str
    source: str
    dest: str
    qty_in: float | None
    qty_out: float | None
    running: float | None
    return_of: str
    lot: str
    note: str
    demand: float | None = None


@dataclass
class SummaryRule:
    label: str
    formula: str
    documents: str
    how: str


@dataclass
class Book:
    path: str
    product_name: str
    sku: str
    lot: str
    store_name: str
    warehouse: str
    location: str
    snapshot: date
    company: str
    quant_id: str
    live_on_hand: float | None
    live_reserved: float | None
    live_free: float | None
    transit_on_hand: float | None
    variant: str = ""
    stock_updated_at: datetime | None = None
    incoming_qty: float | None = None
    moves: list[Move] = field(default_factory=list)
    rules: dict[str, SummaryRule] = field(default_factory=dict)

    @property
    def first_day(self) -> date | None:
        days = [row.when.date() for row in self.moves]
        return min(days) if days else None

    @property
    def last_day(self) -> date | None:
        days = [row.when.date() for row in self.moves]
        return max(days) if days else None


_cache: dict[str, tuple[int, Book]] = {}


def _num(value: object) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _snapshot_date(subtitle: str) -> date:
    match = SNAPSHOT_RE.search(subtitle)
    if not match:
        raise ValueError("Snapshot date is missing from the Summary subtitle.")
    month = MONTHS.get(match.group(2).lower()[:3])
    if not month:
        raise ValueError("Snapshot month could not be read.")
    return date(int(match.group(3)), month, int(match.group(1)))


def load_book(path: str | None = None) -> Book:
    settings = get_settings()
    chosen = (path or settings.reconciliation_workbook or "").strip()
    if not chosen:
        raise FileNotFoundError("RECONCILIATION_WORKBOOK is not set.")
    file_path = Path(chosen)
    if not file_path.is_file():
        raise FileNotFoundError(f"Reconciliation workbook not found: {file_path}")
    stamp = file_path.stat().st_mtime_ns
    cached = _cache.get(str(file_path))
    if cached and cached[0] == stamp:
        return cached[1]
    book = _read(file_path)
    _cache[str(file_path)] = (stamp, book)
    return book


def _read(file_path: Path) -> Book:
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Cannot parse header or footer")
        workbook = load_workbook(file_path, data_only=False, read_only=False)
    try:
        summary = workbook["Summary"]
        title = _text(summary["A1"].value)
        subtitle = _text(summary["A2"].value)
        sku_match = re.search(r"\[(\d+)\]", title)
        name_match = re.search(r"\]\s*(.+?)\s+stock reconciliation", title, re.IGNORECASE)
        rules: dict[str, SummaryRule] = {}
        for row in range(5, 17):
            label = _text(summary.cell(row, 1).value)
            if not label:
                continue
            formula = summary.cell(row, 2).value
            rules[label] = SummaryRule(
                label=label,
                formula=formula if isinstance(formula, str) else "",
                documents=_text(summary.cell(row, 3).value),
                how=_text(summary.cell(row, 4).value),
            )
        demand_by_doc: dict[str, float] = {}
        cancelled = workbook["Cancelled"]
        for row in range(2, cancelled.max_row + 1):
            doc = _text(cancelled.cell(row, 2).value)
            demand = _num(cancelled.cell(row, 9).value)
            if doc:
                demand_by_doc[doc] = demand
        ledger = workbook["Ledger"]
        moves: list[Move] = []
        for row in range(2, ledger.max_row + 1):
            when = ledger.cell(row, 1).value
            if not isinstance(when, datetime):
                continue
            doc = _text(ledger.cell(row, 2).value)
            moves.append(
                Move(
                    when=when,
                    doc=doc,
                    bucket=_text(ledger.cell(row, 3).value),
                    state=_text(ledger.cell(row, 4).value),
                    sale=_text(ledger.cell(row, 5).value),
                    partner=_text(ledger.cell(row, 6).value),
                    source=_text(ledger.cell(row, 7).value),
                    dest=_text(ledger.cell(row, 8).value),
                    qty_in=_num(ledger.cell(row, 9).value),
                    qty_out=_num(ledger.cell(row, 10).value),
                    running=_num(ledger.cell(row, 11).value),
                    return_of=_text(ledger.cell(row, 12).value),
                    lot=_text(ledger.cell(row, 13).value),
                    note=_text(ledger.cell(row, 14).value),
                    demand=demand_by_doc.get(doc),
                )
            )
        moves.sort(key=lambda item: item.when)
        live = workbook["Live Stock"]
        stock_row = None
        transit_qty = None
        for row in range(2, live.max_row + 1):
            location = _text(live.cell(row, 1).value)
            if location.startswith("BL-D/Stock"):
                stock_row = row
            elif "transit" in location.lower():
                transit_qty = _num(live.cell(row, 3).value)
        if stock_row is None:
            raise ValueError("Live Stock sheet has no BL-D/Stock row.")
        warehouse = subtitle.split("|")[0].strip()
        lot_match = re.search(r"Lot\s+(\S+)", subtitle)
        return Book(
            path=str(file_path),
            product_name=(name_match.group(1).strip() if name_match else ""),
            sku=sku_match.group(1) if sku_match else "",
            lot=lot_match.group(1) if lot_match else "",
            store_name="Lavkush Dark Store",
            warehouse=warehouse,
            location="BL-D/Stock",
            snapshot=_snapshot_date(subtitle),
            company=_text(live.cell(stock_row, 8).value),
            quant_id=_text(live.cell(stock_row, 7).value),
            live_on_hand=_num(live.cell(stock_row, 3).value),
            live_reserved=_num(live.cell(stock_row, 4).value),
            live_free=_num(live.cell(stock_row, 5).value),
            transit_on_hand=transit_qty,
            moves=moves,
            rules=rules,
        )
    finally:
        workbook.close()
